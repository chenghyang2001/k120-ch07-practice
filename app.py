"""旗標訂餐系統 Flask 後端：提供餐廳菜單分析（假資料）與訂單寫入 orders.xlsx。"""
import logging
import threading
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS
from openpyxl import Workbook, load_workbook

# 以程式所在目錄為基準，避免從不同工作目錄啟動時找不到檔案，也避免硬編碼使用者路徑
BASE_DIR = Path(__file__).resolve().parent
ORDERS_FILE = BASE_DIR / "orders.xlsx"
HEADERS = ["姓名", "分機", "品項", "選項", "備註", "數量", "金額", "送出時間"]
TIME_FORMAT = "%Y-%m-%d %H:%M:%S"
MAX_QUANTITY = 99
# Excel 會把這些字元開頭的內容當公式執行（CSV/公式注入），寫入前需跳脫
FORMULA_PREFIXES = ("=", "+", "-", "@")

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

app = Flask(__name__)
CORS(app)

# openpyxl 不是 process-safe 的檔案寫入，多個請求同時 load→save 會互相覆蓋，故以鎖序列化
orders_lock = threading.Lock()

# TODO: 之後改為真正爬取/分析餐廳網址的菜單，目前先回傳固定假資料供前端串接
FAKE_MENU_ITEMS = [
    {"id": 1, "name": "大腸麵線", "price": 60, "options": ["正常", "少辣", "不辣"]},
    {"id": 2, "name": "臭豆腐", "price": 70, "options": ["正常", "加辣", "不辣"]},
    {"id": 3, "name": "蚵仔煎", "price": 75, "options": ["正常", "醬多", "醬少"]},
    {"id": 4, "name": "滷肉飯", "price": 40, "options": ["小碗", "大碗"]},
    {"id": 5, "name": "珍珠奶茶", "price": 55, "options": ["正常甜", "半糖", "無糖"]},
]


class OrderValidationError(ValueError):
    """訂單內容不合法時拋出，訊息會直接回給前端（因此只放使用者可讀的中文說明）。"""


def error_response(message, status_code):
    """統一錯誤回應格式。"""
    return jsonify({"error": message}), status_code


def sanitize_cell(value):
    """字串若以公式字元開頭則加單引號，防止 Excel 公式注入；非字串原樣回傳。"""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def is_valid_url(url):
    """檢查是否為 http(s) 開頭的非空字串網址。"""
    if not isinstance(url, str):
        return False
    stripped_url = url.strip()
    return stripped_url.startswith(("http://", "https://"))


def is_real_number(value):
    """bool 是 int 的子類別，需排除以免 True 被當成 1。"""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def require_non_empty_text(value, field_label):
    """驗證欄位為非空字串並回傳 strip 後的值。"""
    if not isinstance(value, str) or not value.strip():
        raise OrderValidationError(f"{field_label}不可為空")
    return value.strip()


def normalize_extension(raw_extension):
    """分機允許字串或整數；整數轉為字串（bool 不接受）。"""
    if isinstance(raw_extension, int) and not isinstance(raw_extension, bool):
        raw_extension = str(raw_extension)
    return require_non_empty_text(raw_extension, "分機")


def normalize_optional_text(value, field_label):
    """選項/備註可省略；缺省或 None 視為空字串，其他非字串型別視為錯誤。"""
    if value is None:
        return ""
    if not isinstance(value, str):
        raise OrderValidationError(f"{field_label}必須是文字")
    return value.strip()


def validate_quantity(quantity, index):
    """數量必須是 1~99 的整數（bool 不算）。"""
    is_int = isinstance(quantity, int) and not isinstance(quantity, bool)
    if not is_int or not 1 <= quantity <= MAX_QUANTITY:
        raise OrderValidationError(f"第 {index} 個品項的數量必須是 1~{MAX_QUANTITY} 的整數")
    return quantity


def validate_price(price, index):
    """金額單價必須是非負數字（bool 不算）。"""
    if not is_real_number(price) or price < 0:
        raise OrderValidationError(f"第 {index} 個品項的價格必須是非負數字")
    return price


def validate_item(raw_item, index):
    """驗證並正規化單一品項，index 從 1 起算以便訊息對使用者友善。"""
    if not isinstance(raw_item, dict):
        raise OrderValidationError(f"第 {index} 個品項格式錯誤")
    return {
        "name": require_non_empty_text(raw_item.get("name"), f"第 {index} 個品項名稱"),
        "option": normalize_optional_text(raw_item.get("option"), f"第 {index} 個品項選項"),
        "note": normalize_optional_text(raw_item.get("note"), f"第 {index} 個品項備註"),
        "quantity": validate_quantity(raw_item.get("quantity"), index),
        "price": validate_price(raw_item.get("price"), index),
    }


def validate_order(payload):
    """驗證整張訂單，成功回傳正規化後的 dict，失敗拋 OrderValidationError。"""
    if not isinstance(payload, dict):
        raise OrderValidationError("請求內容必須是 JSON 物件")
    name = require_non_empty_text(payload.get("name"), "姓名")
    extension = normalize_extension(payload.get("extension"))
    raw_items = payload.get("items")
    if not isinstance(raw_items, list) or not raw_items:
        raise OrderValidationError("至少需要一個品項")
    items = [validate_item(raw_item, index) for index, raw_item in enumerate(raw_items, start=1)]
    return {"name": name, "extension": extension, "items": items}


def build_order_rows(order, submitted_at):
    """把訂單展開成每品項一列；同一張訂單共用同一個送出時間。"""
    rows = []
    for item in order["items"]:
        amount = item["price"] * item["quantity"]
        rows.append([
            sanitize_cell(order["name"]),
            sanitize_cell(order["extension"]),
            sanitize_cell(item["name"]),
            sanitize_cell(item["option"]),
            sanitize_cell(item["note"]),
            item["quantity"],
            amount,
            submitted_at,
        ])
    return rows


def open_orders_workbook():
    """檔案存在就載入，否則建立新活頁簿並寫入標題列。"""
    if ORDERS_FILE.exists():
        return load_workbook(ORDERS_FILE)
    workbook = Workbook()
    workbook.active.title = "訂單"
    workbook.active.append(HEADERS)
    return workbook


def append_order_rows(rows):
    """在鎖保護下把多列附加到 orders.xlsx 並存檔。"""
    with orders_lock:
        workbook = open_orders_workbook()
        try:
            worksheet = workbook.active
            for row in rows:
                worksheet.append(row)
            workbook.save(ORDERS_FILE)
        finally:
            workbook.close()


@app.route("/")
def serve_index():
    """回傳前端頁面，讓使用者可直接開 http://127.0.0.1:5000。"""
    return send_from_directory(BASE_DIR, "index.html")


@app.route("/api/analyze", methods=["POST"])
def analyze_restaurant():
    """分析餐廳網址並回傳菜單品項（目前為假資料）。"""
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return error_response("請求內容必須是 JSON 物件", 400)
    if not is_valid_url(payload.get("url")):
        return error_response("請提供以 http:// 或 https:// 開頭的餐廳網址", 400)
    # TODO: 依 payload["url"] 實際抓取並解析餐廳菜單，取代 FAKE_MENU_ITEMS
    return jsonify({"items": FAKE_MENU_ITEMS}), 200


@app.route("/api/order", methods=["POST"])
def submit_order():
    """接收訂單、驗證後寫入 orders.xlsx。"""
    try:
        order = validate_order(request.get_json(silent=True))
    except OrderValidationError as validation_error:
        return error_response(str(validation_error), 400)

    submitted_at = datetime.now().strftime(TIME_FORMAT)
    rows = build_order_rows(order, submitted_at)
    total_amount = sum(row[6] for row in rows)
    try:
        append_order_rows(rows)
    except PermissionError:
        logger.warning("orders.xlsx 無法寫入（可能被 Excel 開啟）：%s", ORDERS_FILE)
        return error_response("orders.xlsx 被其他程式開啟中，請關閉 Excel 後再試", 503)
    except Exception:  # noqa: BLE001 — 最外層保護，細節只進 log 不回給前端
        logger.exception("寫入訂單時發生未預期錯誤")
        return error_response("伺服器內部錯誤", 500)
    return jsonify({"message": "訂單已送出", "rows": len(rows), "total": total_amount}), 201


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
